"""`excel.py` のシート2「全地点（管理用）」のユニットテスト（設計 6.5 / 6.6 / 13 章）。"""

import io
import zipfile
from datetime import datetime

import openpyxl
import pytest
from openpyxl.utils.cell import range_boundaries
from saas_a_billing.excel import (
    COLOR_EXCLUDED,
    COLOR_FAILED,
    COLOR_MISSING_TEXT,
    COLOR_NO_DATA,
    COLOR_REVIEW_TEXT,
    COLOR_TOTAL_HIGHLIGHT,
    CURRENCY_FORMAT,
    DATETIME_FORMAT,
    DETAIL_COLUMNS,
    FONT_NAME,
    RETENTION_3Y,
    SHEET1_NAME,
    SHEET2_NAME,
    SUMMARY_COLUMNS,
    TOTAL_FONT_SIZE,
    TOTAL_LABEL,
    build_workbook,
    spec_block_lines,
)

from tests.conftest import make_result, make_result_device, make_result_tenant

ALL_DEVICE_CDS = (
    "10000001-01-LoggerA-400",
    "10000002-01-LoggerA-400",
    "10000003-01-LoggerA-400",
    "10000004-01-LoggerA-400",
    "10000005-01-LoggerA-400",
    "20000001-01-CamA-MC",
    "20000002-01-CamA-MC-2",
    "20000003-01-LoggerA-400",
    "30000001-01-LoggerA-400",
    "40000001-01-LoggerA-400",
    "50000001-01-LoggerA-400",
    "10000009-01-LoggerA-400",
)
DEVICE_CD_COLUMN = 4
AMOUNT_COLUMN = 12
HOSTILE_DEVICE_CD = "90000001-01-LoggerA-400"


def argb(color: str) -> str:
    """`#C00000` → openpyxl が返す `FFC00000`。"""
    return f"FF{color.lstrip('#').upper()}"


def load(result: dict) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(build_workbook(result)))


def values(worksheet) -> list[list]:
    return [[cell.value for cell in row] for row in worksheet.iter_rows()]


def at(worksheet, row: int, col: int):
    """0 始まりの行・列でセルを取る。"""
    return worksheet.cell(row=row + 1, column=col + 1)


def find_row(worksheet, text: str, col: int = 0) -> int:
    """`col` の値が `text` と一致する最初の行（0 始まり）。"""
    for index, row in enumerate(values(worksheet)):
        if row[col] == text:
            return index
    raise AssertionError(f"{text!r} の行がありません")


def device_row(worksheet, device_cd: str) -> int:
    return find_row(worksheet, device_cd, col=DEVICE_CD_COLUMN)


def font_color(cell) -> str | None:
    return cell.font.color.rgb if cell.font.color is not None else None


def fill_color(cell) -> str | None:
    return cell.fill.start_color.rgb if cell.fill.patternType else None


def detail_table(worksheet):
    return next(iter(worksheet.tables.values()))


def detail_rows(worksheet) -> list[list]:
    """テーブル範囲のうち見出しを除いたデータ行。"""
    min_col, min_row, max_col, max_row = range_boundaries(detail_table(worksheet).ref)
    return [
        [worksheet.cell(row=row, column=col).value for col in range(min_col, max_col + 1)]
        for row in range(min_row + 1, max_row + 1)
    ]


@pytest.fixture
def result() -> dict:
    return make_result()


@pytest.fixture
def sheet2(result):
    return load(result)[SHEET2_NAME]


def test_sheet2_starts_with_spec_block(result, sheet2):
    lines = spec_block_lines(result)
    assert [at(sheet2, index, 0).value for index in range(len(lines))] == lines


def test_judgement_counts_block_has_six_categories(sheet2):
    start = find_row(sheet2, "判定別件数") + 2
    labels = [at(sheet2, start + offset, 0).value for offset in range(6)]
    assert len(set(labels)) == 6
    for keyword in ("課金", "通信なし", "除外", "要確認", "取得失敗", "消えた地点"):
        assert any(keyword in label for label in labels)
    assert any("判定を問わず" in label for label in labels)


def test_judgement_counts_match_summary(result, sheet2):
    summary = result["summary"]
    start = find_row(sheet2, "判定別件数") + 2
    assert [at(sheet2, start + offset, 1).value for offset in range(6)] == [
        summary["billable_count"],
        summary["no_data_count"],
        summary["excluded_count"],
        summary["review_count"],
        summary["failed_tenant_count"],
        summary["missing_count"],
    ]


def test_tenant_summary_columns(sheet2):
    assert SUMMARY_COLUMNS == (
        "テナントCD",
        "テナント名",
        "保存期間",
        "地点数",
        "課金",
        "通信なし",
        "除外",
        "要確認",
        "取得失敗",
        "小計",
    )
    header = find_row(sheet2, "テナントCD")
    assert [at(sheet2, header, col).value for col in range(len(SUMMARY_COLUMNS))] == list(SUMMARY_COLUMNS)
    assert [at(sheet2, header + offset, 0).value for offset in range(1, 6)] == [
        "TNT001",
        "TNT002",
        "TNT003",
        "TNT004",
        "TNT005",
    ]


def test_tenant_summary_subtotal_matches_devices(sheet2):
    row = find_row(sheet2, "TNT001")
    assert [at(sheet2, row, col).value for col in range(2, 10)] == ["3年", 5, 3, 1, 1, 2, 0, 3450]
    subtotal = at(sheet2, row, 9)
    assert subtotal.number_format == CURRENCY_FORMAT
    assert subtotal.data_type != "f"
    amounts = [row[AMOUNT_COLUMN] for row in detail_rows(sheet2) if row[0] == "TNT001"]
    assert sum(amount for amount in amounts if amount is not None) == 3450


def test_sheet2_spec_block_ends_with_total_amount_row(result, sheet2):
    row = len(spec_block_lines(result))
    label, amount = at(sheet2, row, 0), at(sheet2, row, 1)
    assert label.value == TOTAL_LABEL == "合計金額"
    assert (amount.value, amount.data_type, amount.number_format) == (7200, "n", CURRENCY_FORMAT)
    assert TOTAL_FONT_SIZE >= 14
    assert label.font.bold is True and amount.font.bold is True
    assert label.font.size == amount.font.size == TOTAL_FONT_SIZE
    assert amount.alignment.horizontal == "right"
    assert fill_color(label) == fill_color(amount) == argb(COLOR_TOTAL_HIGHLIGHT)


def test_tenant_summary_ends_with_total_row(sheet2):
    header = find_row(sheet2, "テナントCD")
    total_row = find_row(sheet2, TOTAL_LABEL, col=8)
    assert total_row == header + 6
    amount = at(sheet2, total_row, 9)
    assert amount.value == sum(at(sheet2, header + offset, 9).value for offset in range(1, 6)) == 7200
    assert amount.font.bold is True
    assert amount.number_format == CURRENCY_FORMAT
    assert fill_color(amount) == argb(COLOR_TOTAL_HIGHLIGHT)
    assert fill_color(at(sheet2, total_row, 8)) == argb(COLOR_TOTAL_HIGHLIGHT)


def test_detail_table_has_fifteen_columns(sheet2):
    assert DETAIL_COLUMNS == (
        "テナントCD",
        "テナント名",
        "保存期間",
        "現場",
        "地点CD",
        "地点名",
        "フォーマット",
        "項目数",
        "カメラ",
        "最終受信日時",
        "判定",
        "単価",
        "金額",
        "要確認",
        "備考",
    )
    assert [column.name for column in detail_table(sheet2).tableColumns] == list(DETAIL_COLUMNS)


def test_detail_table_is_registered_as_excel_table_with_autofilter(sheet2):
    assert len(sheet2.tables) == 1
    table = detail_table(sheet2)
    assert table.headerRowCount == 1
    assert table.autoFilter is not None
    assert table.autoFilter.ref == table.ref


def test_no_freeze_panes(sheet2):
    assert sheet2.freeze_panes is None


def test_all_devices_are_listed_including_no_data_and_excluded(sheet2):
    listed = [row[DEVICE_CD_COLUMN] for row in detail_rows(sheet2)]
    assert len(listed) == len(ALL_DEVICE_CDS)
    assert set(listed) == set(ALL_DEVICE_CDS)


def test_no_data_row_is_amber(sheet2):
    row = device_row(sheet2, "10000003-01-LoggerA-400")
    assert at(sheet2, row, 10).value == "通信なし"
    for col in (0, 5, 10):
        assert fill_color(at(sheet2, row, col)) == argb(COLOR_NO_DATA)


def test_excluded_row_is_blue(sheet2):
    row = device_row(sheet2, "10000004-01-LoggerA-400")
    assert at(sheet2, row, 10).value == "除外"
    for col in (0, 5, 10):
        assert fill_color(at(sheet2, row, col)) == argb(COLOR_EXCLUDED)


def test_failed_row_is_gray(sheet2):
    row = device_row(sheet2, "50000001-01-LoggerA-400")
    assert at(sheet2, row, 10).value == "取得失敗"
    for col in (0, 5, 10):
        assert fill_color(at(sheet2, row, col)) == argb(COLOR_FAILED)


def test_missing_devices_are_listed_in_gray_text(sheet2):
    row = device_row(sheet2, "10000009-01-LoggerA-400")
    assert detail_rows(sheet2)[-1][DEVICE_CD_COLUMN] == "10000009-01-LoggerA-400"
    assert at(sheet2, row, 10).value == "消えた地点"
    assert at(sheet2, row, AMOUNT_COLUMN).value is None
    for col in (0, 5, 10):
        assert font_color(at(sheet2, row, col)) == argb(COLOR_MISSING_TEXT)


def test_review_flags_are_red_text_with_reason(sheet2):
    expected = {
        "10000005-01-LoggerA-400": "OVER_MAX_ITEMS",
        "20000003-01-LoggerA-400": "NO_ITEMS",
        "40000001-01-LoggerA-400": "RETENTION_UNKNOWN",
        "10000003-01-LoggerA-400": "NO_ITEMS",
    }
    for device_cd, flag in expected.items():
        row = device_row(sheet2, device_cd)
        reason = at(sheet2, row, 13)
        assert flag in reason.value
        assert font_color(reason) == argb(COLOR_REVIEW_TEXT)
        assert font_color(at(sheet2, row, DEVICE_CD_COLUMN)) == argb(COLOR_REVIEW_TEXT)


def test_retention_cell_color_overrides_row_color(sheet2):
    row = device_row(sheet2, "10000003-01-LoggerA-400")
    retention = at(sheet2, row, 2)
    assert retention.value == "3年"
    assert fill_color(retention) == argb(RETENTION_3Y[0])
    assert font_color(retention) == argb(RETENTION_3Y[1])


def test_camera_column_shows_source(sheet2):
    assert at(sheet2, device_row(sheet2, "20000001-01-CamA-MC"), 8).value == "export"
    assert at(sheet2, device_row(sheet2, "20000002-01-CamA-MC-2"), 8).value == "設定"
    assert at(sheet2, device_row(sheet2, "10000001-01-LoggerA-400"), 8).value == "-"


def test_boundary_only_note_in_remarks(sheet2):
    assert "境界日のみ受信" in at(sheet2, device_row(sheet2, "20000001-01-CamA-MC"), 14).value


def test_device_description_in_remarks(sheet2):
    assert "堤防右岸" in at(sheet2, device_row(sheet2, "10000002-01-LoggerA-400"), 14).value


def test_last_received_at_is_datetime_type_with_format(sheet2):
    received = at(sheet2, device_row(sheet2, "10000002-01-LoggerA-400"), 9)
    assert received.value == datetime(2026, 10, 20, 23, 55)
    assert received.value.tzinfo is None
    assert received.number_format == DATETIME_FORMAT
    assert at(sheet2, device_row(sheet2, "10000003-01-LoggerA-400"), 9).value is None


def test_excluded_unit_price_has_strikethrough(sheet2):
    excluded = at(sheet2, device_row(sheet2, "10000004-01-LoggerA-400"), 11)
    assert excluded.value == 750
    assert excluded.font.strike is True
    assert at(sheet2, device_row(sheet2, "10000001-01-LoggerA-400"), 11).font.strike is not True


def test_amount_zero_for_no_data_and_excluded(sheet2):
    for device_cd in ("10000003-01-LoggerA-400", "10000004-01-LoggerA-400"):
        amount = at(sheet2, device_row(sheet2, device_cd), AMOUNT_COLUMN)
        assert amount.value == 0
        assert amount.number_format == CURRENCY_FORMAT


def test_font_is_meiryo_on_sheet2(sheet2):
    for row in sheet2.iter_rows():
        for cell in row:
            assert cell.font.name == FONT_NAME


def test_detail_headers_remain_without_a_table_when_no_devices(result):
    """全テナントが `device/list` から失敗した RUN でもテーブル無しで開ける（設計 11.1）。"""
    empty = make_result(
        tenants=[],
        missing_devices=[],
        summary=dict.fromkeys(result["summary"], 0),
    )
    worksheet = load(empty)[SHEET2_NAME]
    assert worksheet.tables == {}
    header = find_row(worksheet, "地点CD", col=DEVICE_CD_COLUMN)
    assert [at(worksheet, header, col).value for col in range(len(DETAIL_COLUMNS))] == list(DETAIL_COLUMNS)


def test_build_workbook_returns_bytes_and_opens_in_openpyxl(result):
    data = build_workbook(result)
    assert isinstance(data, bytes)
    assert data[:2] == b"PK"
    assert openpyxl.load_workbook(io.BytesIO(data)).sheetnames == [SHEET1_NAME, SHEET2_NAME]


def hostile_result(**device_overrides) -> dict:
    """外部 API 由来の名前だけを差し替えた最小の `result`（SEC-01）。"""
    return make_result(
        tenants=[
            make_result_tenant(
                tenant_cd="TNT900",
                tenant_name="=SUM(1)",
                device_count_api=1,
                subtotal=750,
                counts={"billable": 1, "no_data": 0, "excluded": 0, "review": 0},
                devices=[make_result_device(device_cd=HOSTILE_DEVICE_CD, **device_overrides)],
            )
        ],
        missing_devices=[],
    )


def test_formula_like_names_are_written_as_text():
    """`=` `+` `@` で始まる名前が数式セルにならない（SEC-01）。"""
    worksheet = load(hostile_result(device_name="+1", group_names=["@x"], format_name="=SUM(1)"))[SHEET2_NAME]
    row = device_row(worksheet, HOSTILE_DEVICE_CD)

    assert [at(worksheet, row, col).value for col in (1, 3, 5, 6)] == ["=SUM(1)", "@x", "+1", "=SUM(1)"]
    assert {at(worksheet, row, col).data_type for col in (1, 3, 5, 6)} == {"s"}


def test_array_formula_like_names_are_written_as_text():
    """`{=...}` も配列数式にしない（SEC-01）。"""
    worksheet = load(hostile_result(device_name="{=SUM(1)}"))[SHEET2_NAME]
    cell = at(worksheet, device_row(worksheet, HOSTILE_DEVICE_CD), 5)

    assert cell.value == "{=SUM(1)}"
    assert cell.data_type == "s"


def test_url_like_names_do_not_become_hyperlinks():
    """`strings_to_urls` を切っているのでリンク化しない（SEC-01）。"""
    worksheet = load(hostile_result(device_name="http://attacker.example/x"))[SHEET2_NAME]
    cell = at(worksheet, device_row(worksheet, HOSTILE_DEVICE_CD), 5)

    assert cell.value == "http://attacker.example/x"
    assert cell.hyperlink is None


def test_workbook_has_no_formula_elements():
    """シート XML に `<f>` 要素が 1 つも無い（SEC-01）。"""
    hostile = hostile_result(device_name='=HYPERLINK("http://attacker.example/","請求書")')

    with zipfile.ZipFile(io.BytesIO(build_workbook(hostile))) as archive:
        sheets = [name for name in archive.namelist() if name.startswith("xl/worksheets/sheet")]
        assert sheets
        for name in sheets:
            assert b"<f>" not in archive.read(name)
