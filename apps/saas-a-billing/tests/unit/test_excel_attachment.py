"""メール添付のシート範囲。XLSX の内部にも管理用データを残さない。"""

import io
import zipfile

import openpyxl
import pytest
from saas_a_billing.excel import MISSING_LABEL, SHEET1_NAME, SHEET2_NAME, build_workbook, footnote_lines

from tests.conftest import make_result, make_result_device, make_result_tenant


@pytest.fixture
def result():
    device = make_result_device()
    unlisted = make_result_device(
        device_cd="unlisted-only",
        device_name="UNLISTED_SECRET",
        judgement="no_data",
        amount=0,
    )
    tenant = make_result_tenant(devices=[device, unlisted])
    return make_result(
        tenants=[tenant],
        missing_devices=[
            {
                "tenant_cd": tenant.get("tenant_cd", "TNT001"),
                "tenant_name": tenant["tenant_name"],
                "device_cd": "gone-only",
                "device_name": "MISSING_SECRET",
                "group_names": [],
                "last_seen_month": "2026-09",
            }
        ],
    )


def load(data):
    return openpyxl.load_workbook(io.BytesIO(data))


@pytest.mark.parametrize("empty", [False, True])
def test_billing_only_has_one_visible_sheet_and_no_management_data(result, empty):
    if empty:
        result["tenants"] = []
    full = load(build_workbook(result))
    data = build_workbook(result, billing_only=True)
    book = load(data)
    assert book.sheetnames == [SHEET1_NAME]
    assert book[SHEET1_NAME].sheet_state == "visible"
    assert full.sheetnames == [SHEET1_NAME, SHEET2_NAME]
    assert MISSING_LABEL in [cell.value for row in full[SHEET2_NAME] for cell in row]
    assert not any(
        MISSING_LABEL in cell.value for row in book[SHEET1_NAME] for cell in row if isinstance(cell.value, str)
    )
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert "xl/worksheets/sheet2.xml" not in archive.namelist()
        xml = "\n".join(archive.read(name).decode() for name in archive.namelist() if name.endswith(".xml"))
    for forbidden in (SHEET2_NAME, MISSING_LABEL, "UNLISTED_SECRET", "MISSING_SECRET", "シート2", "シート 2"):
        assert forbidden not in xml
    assert not book.defined_names


def test_billing_only_sheet1_matches_full_except_footnotes(result):
    full = load(build_workbook(result))[SHEET1_NAME]
    limited = load(build_workbook(result, billing_only=True))[SHEET1_NAME]
    footnote_row = next(cell.row for row in full for cell in row if cell.value == "脚注")
    before = {
        cell.coordinate: cell.value
        for row in full
        for cell in row
        if cell.row <= footnote_row and cell.value is not None
    }
    after = {
        cell.coordinate: cell.value
        for row in limited
        for cell in row
        if cell.row <= footnote_row and cell.value is not None
    }
    assert after == before
    assert limited.merged_cells == full.merged_cells


def test_billing_only_footnotes(result):
    sheet = load(build_workbook(result, billing_only=True))[SHEET1_NAME]
    texts = [cell.value for row in sheet for cell in row if isinstance(cell.value, str)]
    assert "・集計を実行するまでに削除された地点は地点一覧に現れない（金額は計上しない）" in texts
    assert not any("シート2" in text or "境界日のみ受信" in text for text in texts)
    assert any("シート2" in text for text in footnote_lines(result))
