"""`batch.handlers.aggregate` のテスト（設計 6.3 / 6.5 / 6.6 / 6.7 / 6.8 / 6.9 / 7.3 / 11.1）。

moto の DynamoDB / S3 / SQS に実際に読み書きし、`Init` が置いたテナント一覧のスナップ
ショットと `CollectTenant` が置いた parts を入力にして一気通貫で動かす。Excel はモック
せず実物の `excel.build_workbook()` を通し、openpyxl で読み直す（組み上げた
`result.json` が設計 6.6 の Excel 生成に耐えることの証明）。

parts をキー指定で読んでいること（`ListObjectsV2` に依存しないこと）は botocore の
`_make_api_call` を記録して確かめる。Map の `Catch` を通ったテナントは parts が無いまま
`tenant_cd` だけが返るため、そのテナントも取得失敗として集計されることを見る。
"""

import io
import json
import logging
from dataclasses import fields
from datetime import datetime
from types import SimpleNamespace
from unittest.mock import Mock

import boto3
import botocore.client
import openpyxl
import pytest
from batch.handlers import aggregate
from botocore.exceptions import ClientError
from saas_a_billing import billing, messages, repo, storage
from saas_a_billing.period import JST

from example_backoffice import notify
from tests.conftest import make_device, make_price_table, make_settings, make_tenant

VERIFICATION = "verification"
PRODUCTION = "production"
MONTH = "2026-10"
PREVIOUS_MONTH = "2026-09"
ATTEMPT = 1
RUN_TOKEN = "run-token-1"
ACCESS_TOKEN = "dummy-access-token-for-test"
EXPIRES_AT = "2026-10-21 08:48:25"
EXECUTION_ARN = "arn:aws:states:ap-northeast-1:000000000000:execution:saas-a-billing-run:verification-2026-10-1"
SAASA_HOST = "saasa-verification.example.com"
PRICE_EFFECTIVE_FROM = "2025-09-01"
SLACK_CHANNEL = "C0123456789"
SLACK_MENTION = "<!here>"
NOW = datetime(2026, 10, 21, 6, 0, tzinfo=JST)

TNT001 = "TNT001"
TNT002 = "TNT002"
TNT003 = "TNT003"
TNT004 = "TNT004"
TENANT_IDS = {
    TNT001: "00000000-0000-0000-0000-000000000011",
    TNT002: "00000000-0000-0000-0000-000000000012",
    TNT003: "00000000-0000-0000-0000-000000000013",
    TNT004: "00000000-0000-0000-0000-000000000014",
}

WATER = "10000001-01-LoggerA-400"
NO_ITEMS = "10000002-01-LoggerA-400"
DEMO = "10000003-01-LoggerA-400"
CAMERA = "20000001-01-CamA-MC"
FAILED_DEVICE = "30000001-01-LoggerA-400"
GONE = "10000009-01-LoggerA-400"
FAILED_TENANT_GONE = "40000009-01-LoggerA-400"

WATER_PRICE = 650
CAMERA_PRICE = 1400
TOTAL_AMOUNT = WATER_PRICE + CAMERA_PRICE

CLOSING_PERIOD = {
    "from": "2026-09-21T00:00:00+09:00",
    "to": "2026-10-21T00:00:00+09:00",
    "cutoff_day": 20,
    "cutoff_time": "24:00",
    "open": False,
}

COLLECTED = [
    {"tenant_cd": TNT001, "devices": 3, "received": 2, "failed": False},
    {"tenant_cd": TNT002, "devices": 1, "received": 1, "failed": False},
    {"tenant_cd": TNT003, "devices": 1, "received": 0, "failed": True},
    {"tenant_cd": TNT004, "failed": True},
]

DEVICE_FIELDS = (
    "device_cd",
    "device_name",
    "group_cds",
    "group_names",
    "format_name",
    "item_count",
    "device_description",
    "is_camera",
    "camera_source",
    "retention",
    "received",
    "first_received_at",
    "last_received_at",
    "boundary_only",
    "judgement",
    "unit_price",
    "amount",
    "flags",
    "note",
)
RESULT_KEYS = {
    "schema",
    "target_env",
    "month",
    "closing_period",
    "run",
    "settings_snapshot",
    "price_table",
    "summary",
    "tenants",
    "missing_devices",
}
RUN_KEYS = {
    "attempt",
    "trigger",
    "started_at",
    "finished_at",
    "requested_by",
    "env",
    "target_env",
    "saasa_host",
    "execution_arn",
}
TENANT_KEYS = {
    "tenant_id",
    "tenant_cd",
    "tenant_name",
    "tenant_description",
    "tenant_excluded",
    "retention_period",
    "device_count_api",
    "status",
    "error",
    "export_calls",
    "export_bytes",
    "subtotal",
    "counts",
    "devices",
}
EXPECTED_SUMMARY = {
    "tenant_count": 4,
    "device_count": 5,
    "billable_count": 2,
    "no_data_count": 1,
    "excluded_count": 1,
    "review_count": 1,
    "failed_tenant_count": 2,
    "missing_count": 0,
    "total_amount": TOTAL_AMOUNT,
}
METRIC_NAMES = {"RunDuration", "TenantCount", "TenantFailed", "MissingDevices", "TotalAmount"}
LOGGER_NAME = "example_backoffice"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GROUPS = [{"group_cd": "GRP001", "group_name": "現場A"}]


def reset() -> None:
    repo.reset_clients()
    storage.reset_clients()
    notify.reset_clients()


@pytest.fixture
def aws(dynamodb_table, s3_bucket, notify_queue):
    """テーブル・添付バケット・通知キューを立て、boto3 クライアントのキャッシュを捨てる。"""
    reset()
    yield SimpleNamespace(table=dynamodb_table, bucket=s3_bucket, queue=notify_queue)
    reset()


@pytest.fixture
def api_calls(monkeypatch):
    """botocore が投げた API の操作名を順に記録する（一覧走査をしていないことの確認用）。"""
    calls: list[str] = []
    original = botocore.client.BaseClient._make_api_call

    def record(self, operation_name, api_params):
        calls.append(operation_name)
        return original(self, operation_name, api_params)

    monkeypatch.setattr(botocore.client.BaseClient, "_make_api_call", record)
    return calls


def tenant_rows() -> list[dict]:
    """`Init` が S3 に残す `tenant/list` の全属性スナップショット。"""
    return [
        make_tenant(tenant_id=TENANT_IDS[TNT001], tenant_cd=TNT001, device_count=3),
        make_tenant(
            tenant_id=TENANT_IDS[TNT002],
            tenant_cd=TNT002,
            tenant_name="第二テナント",
            retention_period=1,
            device_count=1,
        ),
        make_tenant(tenant_id=TENANT_IDS[TNT003], tenant_cd=TNT003, tenant_name="第三テナント", device_count=1),
        make_tenant(tenant_id=TENANT_IDS[TNT004], tenant_cd=TNT004, tenant_name="第四テナント", device_count=2),
    ]


def make_part(tenant_cd: str, devices: list[dict], **overrides) -> dict:
    """`CollectTenant` が置く `parts/{tenant_cd}.json`（設計 6.4）。"""
    part = {
        "tenant_id": TENANT_IDS[tenant_cd],
        "tenant_cd": tenant_cd,
        "status": "ok",
        "error": None,
        "devices": devices,
        "groups": GROUPS,
        "received": [],
        "camera": [],
        "first_seen": {},
        "last_seen": {},
        "export_calls": 5,
        "export_bytes": 1024,
    }
    part.update(overrides)
    return part


def parts() -> list[dict]:
    """3 テナント分の中間結果。`TNT004` は Map の `Catch` を通ったため parts が無い。"""
    return [
        make_part(
            TNT001,
            [
                make_device(device_cd=WATER),
                make_device(device_cd=NO_ITEMS, device_name="観測局2", format_detail=[]),
                make_device(device_cd=DEMO, device_name="観測局3", device_description="デモ機"),
            ],
            received=[DEMO, WATER],
            first_seen={WATER: "2026/09/21 00:05:00", DEMO: "2026/10/01 09:00:00"},
            last_seen={WATER: "2026/10/20 23:55:00", DEMO: "2026/10/19 09:00:00"},
        ),
        make_part(
            TNT002,
            [
                make_device(
                    device_cd=CAMERA,
                    device_name="カメラ地点1",
                    format_name="CamA-MC",
                    format_detail=[{"item_name": "画像"}],
                    group_cds=["GRP101"],
                    group_names=["現場D"],
                )
            ],
            received=[CAMERA],
            camera=[CAMERA],
            first_seen={CAMERA: "2026/10/20 08:00:00"},
            last_seen={CAMERA: "2026/10/20 08:00:00"},
        ),
        make_part(
            TNT003,
            [make_device(device_cd=FAILED_DEVICE, device_name="観測局9", group_cds=["GRP301"], group_names=["現場F"])],
            status="failed",
            error="HTTPError: 500 Internal Server Error",
            export_calls=1,
            export_bytes=0,
        ),
    ]


def ok_parts() -> list[dict]:
    """取得失敗が 1 件も無い parts。TNT003 を成功にし、Catch を通っていた TNT004 も置く。

    取得失敗があると配布を見送るため（#1006）、配布そのものを見るテストはこちらを使う。
    """
    return [*({**part, "status": "ok", "error": None} for part in parts()), make_part(TNT004, [])]


def put_settings(aws, **overrides) -> dict:
    """配布の検証が主目的なので、メール配布は既定で ON にする（`mail_enabled=False` で切れる）。"""
    settings = make_settings(**{"mail_enabled": True, **overrides})
    aws.table.put_item(Item={**settings, "PK": "SETTINGS", "SK": "current"})
    return settings


def put_price(aws, effective_from: str = PRICE_EFFECTIVE_FROM) -> dict:
    price = make_price_table(effective_from=effective_from)
    aws.table.put_item(Item={**price, "PK": "PRICE", "SK": f"EFFECTIVE#{effective_from}"})
    return price


def prepare_run(target_env: str, *, distribute: bool, attempt: int) -> None:
    """`CollectTenants` が終わった直後の RUN（`running` + トークン保持）を作る。"""
    for _ in range(attempt):
        repo.acquire_run(
            target_env,
            MONTH,
            trigger="scheduled",
            force=True,
            run_token=RUN_TOKEN,
            distribute=distribute,
            stale_before=NOW,
        )
    repo.start_run(target_env, MONTH, attempt, RUN_TOKEN)
    repo.set_execution_arn(target_env, MONTH, attempt, EXECUTION_ARN)
    repo.set_saasa_host(target_env, MONTH, attempt, SAASA_HOST)
    repo.set_token(target_env, MONTH, attempt, ACCESS_TOKEN, EXPIRES_AT)
    repo.set_progress(target_env, MONTH, attempt, len(COLLECTED), len(COLLECTED))


def prepare(
    aws,
    *,
    target_env: str = VERIFICATION,
    distribute: bool = True,
    attempt: int = ATTEMPT,
    all_ok: bool = False,
    **overrides,
):
    """設定・単価表・RUN・テナント一覧スナップショット・parts を用意する。"""
    put_settings(aws, target_env=target_env, **overrides)
    put_price(aws)
    prepare_run(target_env, distribute=distribute, attempt=attempt)
    storage.put_tenants(target_env, MONTH, tenant_rows())
    for part in ok_parts() if all_ok else parts():
        storage.put_part(target_env, MONTH, part["tenant_cd"], part)


def previous_result() -> dict:
    """前月の `result.json`。`GONE` は今回の `device/list` に無い（設計 2.5）。"""
    return {
        "schema": "saas-a-billing.result.v1",
        "month": PREVIOUS_MONTH,
        "tenants": [
            {
                "tenant_cd": TNT001,
                "tenant_name": "サンプル建設",
                "devices": [
                    {"device_cd": WATER, "device_name": "観測局1", "group_names": ["現場A"]},
                    {"device_cd": GONE, "device_name": "撤去済み地点", "group_names": ["現場A"]},
                ],
            }
        ],
    }


def previous_result_with_failed_tenant() -> dict:
    """前月の `result.json` に、今回取得に失敗した TNT004 の地点も入っている状態。"""
    result = previous_result()
    failed_tenant = {
        "tenant_cd": TNT004,
        "tenant_name": "第四テナント",
        "devices": [{"device_cd": FAILED_TENANT_GONE, "device_name": "観測局10", "group_names": ["現場I"]}],
    }
    return {**result, "tenants": [*result["tenants"], failed_tenant]}


def put_previous_run(
    aws, target_env: str = VERIFICATION, status: str = repo.SUCCEEDED, result: dict | None = None
) -> None:
    aws.table.put_item(
        Item={"PK": f"RUN#{target_env}", "SK": PREVIOUS_MONTH, "target_env": target_env, "status": status, "attempt": 1}
    )
    storage.put_result(target_env, PREVIOUS_MONTH, result or previous_result())


def set_distributed_count(aws, count: int, target_env: str = VERIFICATION) -> None:
    """過去の配布回数を持つ RUN にする（#1013）。"""
    aws.table.update_item(
        Key={"PK": f"RUN#{target_env}", "SK": MONTH},
        UpdateExpression="SET #distributed_count = :count",
        ExpressionAttributeNames={"#distributed_count": "distributed_count"},
        ExpressionAttributeValues={":count": count},
    )


def make_event(**overrides) -> dict:
    """`Init` の出力に Map の結果（`ResultPath: $.collected`）が足された入力（設計 6.3）。"""
    event = {
        "run_id": f"{VERIFICATION}-{MONTH}-{ATTEMPT}",
        "target_env": VERIFICATION,
        "month": MONTH,
        "attempt": ATTEMPT,
        "tenants": [{"tenant_id": TENANT_IDS[cd], "tenant_cd": cd} for cd in TENANT_IDS],
        "closing_period": CLOSING_PERIOD,
        "price_effective_from": PRICE_EFFECTIVE_FROM,
        "distribute": True,
        "collected": COLLECTED,
    }
    event.update(overrides)
    return event


def raise_send_error(**kwargs) -> str:
    """SQS への投函が失敗する状況（設計 11.1）。"""
    error = {"Error": {"Code": "AWS.SimpleQueueService.NonExistentQueue", "Message": "queue does not exist"}}
    raise ClientError(error, "SendMessage")


def s3():
    return boto3.client("s3")


def result_of(target_env: str = VERIFICATION, month: str = MONTH) -> dict:
    return storage.get_result(target_env, month)


def run_of(target_env: str = VERIFICATION, month: str = MONTH) -> dict:
    return repo.get_run(target_env, month)


def tenant_of(result: dict, tenant_cd: str) -> dict:
    return next(tenant for tenant in result["tenants"] if tenant["tenant_cd"] == tenant_cd)


def device_of(result: dict, device_cd: str) -> dict:
    return next(
        device for tenant in result["tenants"] for device in tenant["devices"] if device["device_cd"] == device_cd
    )


def queued_messages(aws) -> list[dict]:
    response = boto3.client("sqs").receive_message(QueueUrl=aws.queue, MaxNumberOfMessages=10)
    return [json.loads(message["Body"]) for message in response.get("Messages", [])]


def xlsx_keys(aws) -> list[str]:
    listed = s3().list_objects_v2(Bucket=aws.bucket)
    return [obj["Key"] for obj in listed.get("Contents", []) if obj["Key"].endswith(".xlsx")]


def emf_documents(captured: str) -> list[dict]:
    """標準出力に流れた Powertools Metrics の EMF ドキュメント。"""
    documents = []
    for line in captured.splitlines():
        if not line.startswith("{"):
            continue
        try:
            parsed = json.loads(line)
        except ValueError:
            continue
        if "_aws" in parsed:
            documents.append(parsed)
    return documents


def emitted_names(documents: list[dict]) -> set[str]:
    groups = [group for document in documents for group in document["_aws"]["CloudWatchMetrics"]]
    return {metric["Name"] for group in groups for metric in group["Metrics"]}


def metric_value(document: dict, name: str):
    value = document[name]
    return value[0] if isinstance(value, list) else value


def test_reads_parts_by_key_not_by_listing(aws, api_calls):
    prepare(aws)

    aggregate.handler(make_event(), None)

    assert "ListObjectsV2" not in api_calls
    assert "ListObjects" not in api_calls
    assert api_calls.count("GetObject") >= len(parts())


def test_result_json_matches_schema(aws):
    prepare(aws)

    aggregate.handler(make_event(), None)

    result = result_of()
    assert set(result) == RESULT_KEYS
    assert result["schema"] == "saas-a-billing.result.v1"
    assert result["target_env"] == VERIFICATION
    assert result["month"] == MONTH
    assert result["closing_period"] == CLOSING_PERIOD
    assert set(result["run"]) == RUN_KEYS
    assert result["run"]["saasa_host"] == SAASA_HOST
    assert result["run"]["execution_arn"] == EXECUTION_ARN
    assert all(set(tenant) == TENANT_KEYS for tenant in result["tenants"])


def test_result_json_devices_have_all_fields(aws):
    prepare(aws)

    aggregate.handler(make_event(), None)

    assert len(DEVICE_FIELDS) == 19
    assert set(DEVICE_FIELDS) == {field.name for field in fields(billing.DeviceJudgement)}
    devices = [device for tenant in result_of()["tenants"] for device in tenant["devices"]]
    assert devices and all(set(device) == set(DEVICE_FIELDS) for device in devices)
    water = device_of(result_of(), WATER)
    assert water["judgement"] == "billable"
    assert water["unit_price"] == WATER_PRICE
    assert water["amount"] == WATER_PRICE
    assert water["retention"] == 3


def test_tenant_excluded_marks_tenants_whose_description_matches(aws):
    """テナント説明が除外キーワードに一致するテナントは、地点を見ずに丸ごとデモ扱いにする。"""
    prepare(aws)
    rows = [
        {**row, "tenant_description": "デモ用テナント"} if row["tenant_cd"] == TNT002 else row for row in tenant_rows()
    ]
    storage.put_tenants(VERIFICATION, MONTH, rows)

    aggregate.handler(make_event(), None)

    excluded = {tenant["tenant_cd"]: tenant["tenant_excluded"] for tenant in result_of()["tenants"]}
    assert excluded[TNT002] is True
    assert excluded[TNT001] is False


def test_summary_counts_match_devices(aws):
    prepare(aws)

    output = aggregate.handler(make_event(), None)

    assert output["summary"] == EXPECTED_SUMMARY
    result = result_of()
    assert result["summary"] == EXPECTED_SUMMARY
    devices = [device for tenant in result["tenants"] for device in tenant["devices"]]
    assert len(devices) == EXPECTED_SUMMARY["device_count"]
    assert sum(device["amount"] for device in devices) == EXPECTED_SUMMARY["total_amount"]
    assert device_of(result, NO_ITEMS)["flags"] == ["NO_ITEMS"]
    assert device_of(result, NO_ITEMS)["judgement"] == "no_data"
    assert device_of(result, DEMO)["judgement"] == "excluded"
    assert device_of(result, CAMERA)["boundary_only"] is True


def test_missing_devices_from_previous_month_same_target_env(aws):
    prepare(aws)
    put_previous_run(aws)

    aggregate.handler(make_event(), None)

    result = result_of()
    assert result["missing_devices"] == [
        {
            "tenant_cd": TNT001,
            "tenant_name": "サンプル建設",
            "device_cd": GONE,
            "device_name": "撤去済み地点",
            "group_names": ["現場A"],
            "last_seen_month": PREVIOUS_MONTH,
        }
    ]
    assert result["summary"]["missing_count"] == 1
    assert result["summary"]["total_amount"] == TOTAL_AMOUNT


def test_missing_devices_does_not_cross_target_env(aws, monkeypatch):
    monkeypatch.setenv("TARGET_ENV", PRODUCTION)
    prepare(aws, target_env=PRODUCTION)
    put_previous_run(aws, VERIFICATION)

    aggregate.handler(make_event(target_env=PRODUCTION), None)

    assert result_of(PRODUCTION)["missing_devices"] == []
    assert result_of(PRODUCTION)["summary"]["missing_count"] == 0


def test_no_previous_run_yields_empty_missing_devices(aws):
    prepare(aws)
    storage.put_result(VERIFICATION, PREVIOUS_MONTH, previous_result())

    aggregate.handler(make_event(), None)

    assert result_of()["missing_devices"] == []


def test_failed_tenant_devices_are_not_reported_as_missing(aws):
    """取得失敗テナントの地点は「消えた地点」に出さない（撤去と誤読させない）。"""
    prepare(aws)
    put_previous_run(aws, result=previous_result_with_failed_tenant())

    aggregate.handler(make_event(), None)

    result = result_of()
    assert [device["device_cd"] for device in result["missing_devices"]] == [GONE]
    assert result["summary"]["missing_count"] == 1


def test_missing_tenants_snapshot_is_warned(aws, caplog):
    """テナント一覧のスナップショットが無いと除外判定と保存期間が既定値になるため警告する。"""
    prepare(aws)
    s3().delete_object(Bucket=aws.bucket, Key=storage.tenants_key(VERIFICATION, MONTH))

    with caplog.at_level(logging.WARNING, logger=LOGGER_NAME):
        aggregate.handler(make_event(), None)

    assert "テナント一覧のスナップショット" in caplog.text


def test_failed_tenant_is_counted_and_devices_marked_failed(aws):
    prepare(aws)

    output = aggregate.handler(make_event(), None)

    result = result_of()
    collected = tenant_of(result, TNT003)
    assert collected["status"] == "failed"
    assert collected["error"] == "HTTPError: 500 Internal Server Error"
    assert [device["judgement"] for device in collected["devices"]] == ["failed"]
    assert device_of(result, FAILED_DEVICE)["amount"] == 0
    caught = tenant_of(result, TNT004)
    assert caught["status"] == "failed"
    assert caught["devices"] == []
    assert caught["device_count_api"] == 2
    assert output["summary"]["failed_tenant_count"] == 2
    assert output["summary"]["total_amount"] == TOTAL_AMOUNT


def test_xlsx_is_written_to_s3_with_correct_key(aws):
    prepare(aws)

    output = aggregate.handler(make_event(), None)

    key = storage.xlsx_key(VERIFICATION, MONTH)
    assert output["s3_xlsx_key"] == key
    assert key == f"saas-a-billing/{VERIFICATION}/2026/10/saas_a_billing_2026-10_verification.xlsx"
    stored = s3().get_object(Bucket=aws.bucket, Key=key)
    assert stored["ContentType"] == XLSX_CONTENT_TYPE
    workbook = openpyxl.load_workbook(io.BytesIO(stored["Body"].read()))
    assert workbook.sheetnames == ["請求照合", "全地点（管理用）"]
    cells = [cell.value for row in workbook["全地点（管理用）"].iter_rows() for cell in row]
    texts = [value for value in cells if isinstance(value, str)]
    assert WATER in texts
    assert "第三テナント" in texts


def test_result_json_is_written_to_s3(aws):
    prepare(aws)

    output = aggregate.handler(make_event(), None)

    key = storage.result_key(VERIFICATION, MONTH)
    assert output["s3_result_key"] == key
    stored = s3().get_object(Bucket=aws.bucket, Key=key)
    assert stored["ContentType"] == "application/json"
    body = stored["Body"].read()
    assert "サンプル建設".encode() in body
    assert json.loads(body.decode("utf-8")) == result_of()


def test_run_becomes_succeeded_with_summary_and_keys(aws):
    prepare(aws, distribute=False)

    output = aggregate.handler(make_event(distribute=False), None)

    run = run_of()
    assert run["status"] == repo.SUCCEEDED
    assert run["summary"] == output["summary"]
    assert run["s3_xlsx_key"] == storage.xlsx_key(VERIFICATION, MONTH)
    assert run["s3_result_key"] == storage.result_key(VERIFICATION, MONTH)
    assert run["closing_period"] == CLOSING_PERIOD
    assert run["price_effective_from"] == PRICE_EFFECTIVE_FROM
    assert set(run["metrics"]) == {"duration_sec", "export_calls", "export_bytes"}
    assert run["metrics"]["export_calls"] == 11
    assert run["metrics"]["export_bytes"] == 2048


def test_token_attributes_are_removed_from_run(aws):
    prepare(aws)

    aggregate.handler(make_event(), None)

    run = run_of()
    assert "api_token" not in run
    assert "token_expires_at" not in run


def test_distribute_true_sends_notify_with_attachment(aws):
    prepare(aws, all_ok=True, slack_channel_id=SLACK_CHANNEL, slack_enabled=True, mail_attachment_sheets="all")

    output = aggregate.handler(make_event(), None)

    sent = queued_messages(aws)
    assert len(sent) == 1
    message = sent[0]
    assert message["app"] == "saas-a-billing"
    assert message["channels"] == ["email", "slack"]
    assert message["email"]["to"] == ["devA@example.co.jp"]
    attachment = message["attachments"][0]
    assert attachment["key"] == storage.xlsx_key(VERIFICATION, MONTH)
    assert attachment["filename"] == messages.attachment_filename(MONTH, VERIFICATION)
    assert attachment["bucket"] == aws.bucket
    assert attachment["content_type"] == XLSX_CONTENT_TYPE
    assert output["distributed"] is True


def test_attachment_reuses_existing_key_without_reupload(aws):
    prepare(aws, all_ok=True, mail_attachment_sheets="all")

    aggregate.handler(make_event(), None)

    assert xlsx_keys(aws) == [storage.xlsx_key(VERIFICATION, MONTH)]
    assert queued_messages(aws)[0]["attachments"][0]["key"] == storage.xlsx_key(VERIFICATION, MONTH)


def test_slack_mention_is_sent_as_field_not_in_text(aws):
    prepare(aws, all_ok=True, slack_channel_id=SLACK_CHANNEL, slack_enabled=True, slack_mention=SLACK_MENTION)

    aggregate.handler(make_event(), None)

    slack = queued_messages(aws)[0]["slack"]
    assert slack["channel"] == SLACK_CHANNEL
    assert slack["mention"] == SLACK_MENTION
    assert SLACK_MENTION not in slack["text"]


def test_failed_tenant_skips_distribution_and_notifies(aws):
    """取得失敗があると金額が過少になるため Excel は配らず、再実行を促す（#1006）。"""
    prepare(aws, slack_channel_id=SLACK_CHANNEL, slack_enabled=True)

    output = aggregate.handler(make_event(), None)

    message = queued_messages(aws)[0]
    assert "attachments" not in message
    assert "配布を見送りました" in message["email"]["text"]
    assert "再実行" in message["email"]["text"]
    assert TNT003 in message["email"]["text"]
    assert TNT004 in message["email"]["text"]
    assert f"/saas-a-billing/runs/{VERIFICATION}/{MONTH}" in message["email"]["text"]
    assert output["distributed"] is False


def test_failed_tenant_keeps_run_succeeded_with_note(aws):
    prepare(aws, slack_channel_id=SLACK_CHANNEL, slack_enabled=True)

    aggregate.handler(make_event(), None)

    run = run_of()
    assert run["status"] == repo.SUCCEEDED
    assert run["distribute_note"] == "取得失敗 2 テナントのため配布見送り"
    assert run["notify_message_id"] == queued_messages(aws)[0]["message_id"]
    assert run["notify_channels"] == ["email", "slack"]
    assert run["notified_at"]
    assert "error" not in run


def test_failed_tenant_slack_does_not_use_the_template(aws):
    """配布見送りは定型の連絡なので、設定の文面テンプレートは使わない（設計 6.8）。"""
    prepare(
        aws,
        slack_channel_id=SLACK_CHANNEL,
        slack_enabled=True,
        slack_mention=SLACK_MENTION,
        slack_template="配布しました {total}",
    )

    aggregate.handler(make_event(), None)

    slack = queued_messages(aws)[0]["slack"]
    assert "配布しました" not in slack["text"]
    assert "配布を見送りました" in slack["text"]
    assert slack["mention"] == SLACK_MENTION
    assert SLACK_MENTION not in slack["text"]


def test_failed_tenant_notify_failure_records_error_without_note(aws, monkeypatch):
    prepare(aws, slack_channel_id=SLACK_CHANNEL, slack_enabled=True)
    monkeypatch.setattr(aggregate, "send", raise_send_error)

    output = aggregate.handler(make_event(), None)

    run = run_of()
    assert run["status"] == repo.SUCCEEDED
    assert "NonExistentQueue" in run["error"]
    assert "distribute_note" not in run
    assert output["distributed"] is False


def test_failed_tenant_with_no_distribution_target_keeps_the_existing_note(aws):
    """配布先が両方 OFF のときは取得失敗の有無によらず従来どおりの理由を残す。"""
    prepare(aws, mail_enabled=False, slack_enabled=False)

    aggregate.handler(make_event(), None)

    assert queued_messages(aws) == []
    assert run_of()["distribute_note"] == aggregate.NO_DISTRIBUTION_TARGET


def test_failed_tenant_does_not_notify_when_distribute_is_false(aws):
    prepare(aws, slack_channel_id=SLACK_CHANNEL, slack_enabled=True, distribute=False)

    aggregate.handler(make_event(distribute=False), None)

    assert queued_messages(aws) == []
    assert "distribute_note" not in run_of()


def test_distribute_false_does_not_notify(aws):
    prepare(aws, slack_channel_id=SLACK_CHANNEL, slack_enabled=True, distribute=False)

    output = aggregate.handler(make_event(distribute=False), None)

    assert queued_messages(aws) == []
    assert run_of()["status"] == repo.SUCCEEDED
    assert output["distributed"] is False


def test_run_becomes_distributed_after_notify(aws):
    prepare(aws, all_ok=True, slack_channel_id=SLACK_CHANNEL, slack_enabled=True)

    aggregate.handler(make_event(), None)

    run = run_of()
    assert run["status"] == repo.DISTRIBUTED
    assert run["notify_message_id"] == queued_messages(aws)[0]["message_id"]
    assert run["notify_channels"] == ["email", "slack"]
    assert run["distribute_note"] == ""
    assert run["notified_at"]


def test_mail_only_sends_email_channel(aws):
    """Slack 配布が OFF なら Slack は投函先から外れ、意図的な OFF なので note も出さない。"""
    prepare(aws, all_ok=True, slack_channel_id=SLACK_CHANNEL, slack_enabled=False)

    aggregate.handler(make_event(), None)

    message = queued_messages(aws)[0]
    assert message["channels"] == ["email"]
    assert "slack" not in message
    run = run_of()
    assert run["status"] == repo.DISTRIBUTED
    assert run["notify_channels"] == ["email"]
    assert run["distribute_note"] == ""


def test_slack_only_sends_slack_channel(aws):
    """メール配布が OFF なら宛先が空でも Slack だけに投函する。"""
    prepare(aws, all_ok=True, mail_enabled=False, slack_enabled=True, slack_channel_id=SLACK_CHANNEL, mail_to=[])

    output = aggregate.handler(make_event(), None)

    message = queued_messages(aws)[0]
    assert message["channels"] == ["slack"]
    assert "email" not in message
    assert message["slack"]["channel"] == SLACK_CHANNEL
    run = run_of()
    assert run["status"] == repo.DISTRIBUTED
    assert run["notify_channels"] == ["slack"]
    assert run["distribute_note"] == ""
    assert output["distributed"] is True


def test_both_channels_enabled_sends_both(aws):
    prepare(aws, all_ok=True, slack_channel_id=SLACK_CHANNEL, slack_enabled=True)

    aggregate.handler(make_event(), None)

    assert queued_messages(aws)[0]["channels"] == ["email", "slack"]
    assert run_of()["notify_channels"] == ["email", "slack"]


def test_both_channels_disabled_skips_send_and_keeps_succeeded(aws, monkeypatch):
    """distribute 指定でも配布先がゼロなら送付せず、RUN は succeeded のまま理由を残す。"""
    prepare(aws, mail_enabled=False, slack_enabled=False, slack_channel_id=SLACK_CHANNEL)
    monkeypatch.setattr(aggregate, "send", lambda **_kwargs: pytest.fail("send() must not be called"))

    output = aggregate.handler(make_event(), None)

    assert queued_messages(aws) == []
    run = run_of()
    assert run["status"] == repo.SUCCEEDED
    assert run["distribute_note"] == aggregate.NO_DISTRIBUTION_TARGET
    assert "error" not in run
    assert output["distributed"] is False


def test_distribute_false_does_not_record_distribute_note(aws):
    prepare(aws, mail_enabled=False, slack_enabled=False, distribute=False)

    aggregate.handler(make_event(distribute=False), None)

    assert "distribute_note" not in run_of()


def test_verification_subject_has_suffix_and_filename_has_verification(aws):
    prepare(aws, all_ok=True, slack_channel_id=SLACK_CHANNEL, slack_enabled=True)

    aggregate.handler(make_event(), None)

    message = queued_messages(aws)[0]
    assert message["email"]["subject"].endswith("[検証環境]")
    assert message["slack"]["text"].startswith("【検証環境】")
    assert message["attachments"][0]["filename"] == "saas_a_billing_2026-10_verification.xlsx"


def test_production_has_no_verification_markers(aws, monkeypatch):
    monkeypatch.setenv("TARGET_ENV", PRODUCTION)
    prepare(aws, all_ok=True, target_env=PRODUCTION, slack_channel_id=SLACK_CHANNEL, slack_enabled=True)

    aggregate.handler(make_event(target_env=PRODUCTION), None)

    message = queued_messages(aws)[0]
    assert "[検証環境]" not in message["email"]["subject"]
    assert "【検証環境】" not in message["slack"]["text"]
    assert message["attachments"][0]["filename"] == "saas_a_billing_2026-10.xlsx"


def test_ignores_stale_target_env_in_execution_input(aws, monkeypatch):
    """実行入力に古い `target_env` が残っていても、集計は `TARGET_ENV` 側で行う（計画 R1 / SP-01）。"""
    monkeypatch.setenv("TARGET_ENV", VERIFICATION)
    prepare(aws, all_ok=True, slack_channel_id=SLACK_CHANNEL, slack_enabled=True)

    aggregate.handler(make_event(target_env=PRODUCTION), None)

    assert result_of(VERIFICATION)["target_env"] == VERIFICATION
    assert run_of(VERIFICATION)["status"] == repo.DISTRIBUTED
    assert storage.get_result(PRODUCTION, MONTH) is None
    message = queued_messages(aws)[0]
    assert message["email"]["subject"].endswith("[検証環境]")


def test_notify_failure_keeps_succeeded_and_records_error(aws, monkeypatch):
    prepare(aws, all_ok=True, slack_channel_id=SLACK_CHANNEL, slack_enabled=True)
    monkeypatch.setattr(aggregate, "send", raise_send_error)

    output = aggregate.handler(make_event(), None)

    run = run_of()
    assert run["status"] == repo.SUCCEEDED
    assert "NonExistentQueue" in run["error"]
    assert output["distributed"] is False
    assert queued_messages(aws) == []


def test_settings_snapshot_recorded_in_result(aws):
    prepare(aws, exclude_keywords=["デモ", "試験"], camera_format_names=["CamA-MC-2"])

    aggregate.handler(make_event(), None)

    assert result_of()["settings_snapshot"] == {
        "exclude_keywords": ["デモ", "試験"],
        "camera_format_names": ["CamA-MC-2"],
    }


def test_price_table_recorded_in_result(aws):
    prepare(aws)

    aggregate.handler(make_event(), None)

    price_table = make_price_table()
    assert result_of()["price_table"] == {
        "effective_from": PRICE_EFFECTIVE_FROM,
        "bands": price_table["bands"],
        "camera": price_table["camera"],
        "note": price_table["note"],
    }


def test_metrics_are_emitted(aws, capsys):
    prepare(aws)

    aggregate.handler(make_event(), None)

    documents = emf_documents(capsys.readouterr().out)
    assert METRIC_NAMES.issubset(emitted_names(documents))
    assert documents[0]["_aws"]["CloudWatchMetrics"][0]["Namespace"] == "example-backoffice/saas-a-billing"
    assert metric_value(documents[0], "TenantCount") == EXPECTED_SUMMARY["tenant_count"]
    assert metric_value(documents[0], "TenantFailed") == EXPECTED_SUMMARY["failed_tenant_count"]
    assert metric_value(documents[0], "TotalAmount") == TOTAL_AMOUNT


def test_closing_period_open_true_for_unfinished_month(aws):
    prepare(aws)
    unfinished = {**CLOSING_PERIOD, "open": True}

    aggregate.handler(make_event(closing_period=unfinished), None)

    assert result_of()["closing_period"] == unfinished
    assert run_of()["closing_period"]["open"] is True


def test_first_distribution_has_no_resend_note(aws):
    """再実行でも配るのが初回なら再送ではない（根拠は attempt ではなく配布回数。#1013）。"""
    prepare(aws, all_ok=True, attempt=2)

    aggregate.handler(make_event(attempt=2), None)

    message = queued_messages(aws)[0]
    assert "再送" not in message["email"]["subject"]
    assert result_of()["run"]["attempt"] == 2
    assert run_of()["distributed_count"] == 1


def test_second_distribution_marks_resend_in_subject(aws):
    """2 回目の配布から件名に再送表記を付ける（#1013）。"""
    prepare(aws, all_ok=True, attempt=2)
    set_distributed_count(aws, 1)

    aggregate.handler(make_event(attempt=2), None)

    assert "再送（第 2 回）" in queued_messages(aws)[0]["email"]["subject"]
    assert run_of()["distributed_count"] == 2


def test_partial_failure_subject_uses_the_distribution_count(aws):
    """配布見送りの連絡も同じ根拠で再送表記を付ける。ただし配布回数は進めない（#1013）。"""
    prepare(aws, attempt=2)
    set_distributed_count(aws, 1)

    aggregate.handler(make_event(attempt=2), None)

    assert "再送（第 2 回）" in queued_messages(aws)[0]["email"]["subject"]
    assert run_of()["distributed_count"] == 1


def test_distribute_false_does_not_count_the_distribution(aws):
    """配らなかった実行は配布回数を進めない（#1013）。"""
    prepare(aws, all_ok=True, distribute=False)

    aggregate.handler(make_event(distribute=False), None)

    assert "distributed_count" not in run_of()


def _prepare_attachment_run(aws, *, distribute=True, failed=False, **settings):
    prepare(aws, distribute=distribute, all_ok=not failed, **settings)


def _attachment_book(aws, key):
    obj = boto3.client("s3").get_object(Bucket=aws.bucket, Key=key)
    return openpyxl.load_workbook(io.BytesIO(obj["Body"].read()))


def _month_xlsx_keys(aws):
    objects = boto3.client("s3").list_objects_v2(Bucket=aws.bucket, Prefix=storage.prefix(VERIFICATION, MONTH))
    return {obj["Key"] for obj in objects.get("Contents", []) if obj["Key"].endswith(".xlsx")}


@pytest.mark.parametrize("value", [None, "billing_only", "all"])
def test_mail_attachment_sheets_selects_contents_without_changing_download(aws, caplog, value):
    _prepare_attachment_run(aws, **({"mail_attachment_sheets": value} if value else {}))
    with caplog.at_level(logging.INFO, logger="example_backoffice"):
        aggregate.handler(make_event(distribute=True), None)
    run = run_of()
    full_key = run["s3_xlsx_key"]
    attachment = queued_messages(aws)[0]["attachments"][0]
    assert full_key == storage.xlsx_key(VERIFICATION, MONTH)
    assert _attachment_book(aws, full_key).sheetnames == ["請求照合", "全地点（管理用）"]
    assert attachment["filename"] == full_key.rsplit("/", 1)[1]
    selected = value or "billing_only"
    if selected == "all":
        assert attachment["key"] == full_key
        assert _month_xlsx_keys(aws) == {full_key}
    else:
        assert attachment["key"] == f"{storage.prefix(VERIFICATION, MONTH)}/mail/{attachment['filename']}"
        assert _attachment_book(aws, attachment["key"]).sheetnames == ["請求照合"]
        assert _month_xlsx_keys(aws) == {full_key, attachment["key"]}
    assert "mail_attachment_sheets" not in run
    assert "mail_attachment_sheets" not in result_of()
    record = next(record for record in caplog.records if record.getMessage().startswith("mail attachment selected "))
    details = json.loads(record.getMessage().removeprefix("mail attachment selected "))
    assert details["mail_attachment_sheets"] == selected
    assert details["key"] == attachment["key"]


@pytest.mark.parametrize("value", [None, "", "billing-only", "BILLING_ONLY", 123])
def test_invalid_mail_attachment_sheets_falls_back_to_billing_only(aws, caplog, value):
    _prepare_attachment_run(aws, mail_attachment_sheets=value)
    with caplog.at_level(logging.INFO, logger="example_backoffice"):
        aggregate.handler(make_event(distribute=True), None)

    full_key = run_of()["s3_xlsx_key"]
    message = queued_messages(aws)[0]
    attachment = message["attachments"][0]
    assert attachment["key"] == f"{storage.prefix(VERIFICATION, MONTH)}/mail/{attachment['filename']}"
    assert attachment["key"] != full_key
    assert _attachment_book(aws, attachment["key"]).sheetnames == ["請求照合"]
    assert len(_attachment_book(aws, full_key).sheetnames) == 2
    assert _month_xlsx_keys(aws) == {full_key, attachment["key"]}
    record = next(record for record in caplog.records if record.getMessage().startswith("mail attachment selected "))
    details = json.loads(record.getMessage().removeprefix("mail attachment selected "))
    assert details["mail_attachment_sheets"] == "billing_only"
    assert details["key"] == attachment["key"]


@pytest.mark.parametrize("stage", ["build", "save"])
def test_billing_only_attachment_failure_keeps_succeeded_without_sending(aws, monkeypatch, stage):
    _prepare_attachment_run(aws, mail_attachment_sheets="billing_only")
    send = Mock(wraps=aggregate.send)
    monkeypatch.setattr(aggregate, "send", send)
    owner, name = (aggregate.excel, "build_workbook") if stage == "build" else (storage, "put_xlsx")
    original = getattr(owner, name)
    calls = []
    error_message = f"billing-only workbook {stage} failed"

    def fail_billing_only(*args, billing_only=False):
        calls.append(billing_only)
        if billing_only:
            raise RuntimeError(error_message)
        return original(*args, billing_only=billing_only)

    monkeypatch.setattr(owner, name, fail_billing_only)
    output = aggregate.handler(make_event(distribute=True), None)

    run = run_of()
    assert calls == [False, True]
    send.assert_not_called()
    assert queued_messages(aws) == []
    assert output["distributed"] is False
    assert run["status"] == repo.SUCCEEDED
    assert error_message in run["error"]
    assert "notify_message_id" not in run
    assert "distributed_count" not in run
    full_key = storage.xlsx_key(VERIFICATION, MONTH)
    assert run["s3_xlsx_key"] == full_key
    assert len(_attachment_book(aws, full_key).sheetnames) == 2
    assert _month_xlsx_keys(aws) == {full_key}


@pytest.mark.parametrize("scenario", ["no_distribution", "no_targets", "partial_failure", "slack_only"])
def test_mail_attachment_sheets_does_not_build_unused_version(aws, monkeypatch, scenario):
    _prepare_attachment_run(
        aws,
        distribute=scenario != "no_distribution",
        failed=scenario == "partial_failure",
        mail_enabled=scenario not in ("no_targets", "slack_only"),
        slack_enabled=scenario == "slack_only",
        slack_channel_id="C0123456789",
        mail_attachment_sheets="billing_only",
    )
    builds = []
    original = aggregate.excel.build_workbook

    def record(result, **kwargs):
        builds.append(kwargs)
        return original(result, **kwargs)

    monkeypatch.setattr(aggregate.excel, "build_workbook", record)
    aggregate.handler(make_event(distribute=scenario != "no_distribution"), None)
    assert builds == [{}]
    assert _month_xlsx_keys(aws) == {storage.xlsx_key(VERIFICATION, MONTH)}
    messages = queued_messages(aws)
    if scenario in ("no_distribution", "no_targets"):
        assert messages == []
    elif scenario == "partial_failure":
        assert not messages[0].get("attachments")
    else:
        assert messages[0]["channels"] == ["slack"]
        assert messages[0]["attachments"][0]["key"] == storage.xlsx_key(VERIFICATION, MONTH)


def test_mail_attachment_sheets_does_not_change_slack_payload(aws, monkeypatch):
    _prepare_attachment_run(aws, slack_enabled=True, slack_channel_id="C0123456789", slack_mention="<!here>")
    aggregate.handler(make_event(distribute=False), None)
    result = result_of()
    closing = aggregate._closing_period(make_event())
    settings = repo.get_settings()
    captured = []
    monkeypatch.setattr(aggregate, "send", lambda **payload: captured.append(payload) or "message-id")
    for selected in ("billing_only", "all"):
        aggregate._send_report(
            {**settings, "mail_attachment_sheets": selected},
            result,
            closing,
            run_of()["s3_xlsx_key"],
            ["email", "slack"],
            1,
        )
    assert captured[0]["slack"] == captured[1]["slack"]
    assert captured[0]["email"] == captured[1]["email"]
    assert captured[0]["meta"] == captured[1]["meta"]
    assert captured[0]["channels"] == captured[1]["channels"]
